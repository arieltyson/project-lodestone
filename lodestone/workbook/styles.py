"""Shared formatting for the Excel model, following banking colour conventions.

Blue text = hardcoded input, black = formula, green = link to another sheet.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

FONT = "Arial"
NAVY = "1F2A44"

INPUT = Font(name=FONT, size=10, color="0000FF")
FORMULA = Font(name=FONT, size=10, color="000000")
LINK = Font(name=FONT, size=10, color="008000")
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color=NAVY)
SUBTITLE = Font(name=FONT, size=10, italic=True, color="595959")
HEADER = Font(name=FONT, size=10, bold=True, color="FFFFFF")
NOTE = Font(name=FONT, size=9, italic=True, color="595959")

HEADER_FILL = PatternFill("solid", fgColor=NAVY)
SECTION_FILL = PatternFill("solid", fgColor="DCE3EE")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
TOTAL_BORDER = Border(top=Side(style="thin"))

USD_M = '#,##0.0;(#,##0.0);"-"'
USD_0 = '#,##0;(#,##0);"-"'
USD_2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT_2 = '0.00%;(0.00%);"-"'
MULT = '0.00"x";(0.00"x");"-"'
NUM_3 = '0.000;(0.000);"-"'
YEAR = "0"


def title(ws, text: str, subtitle: str) -> None:
    ws["B1"] = text
    ws["B1"].font = TITLE
    ws["B2"] = subtitle
    ws["B2"].font = SUBTITLE
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2


def section(ws, row: int, text: str, last_col: int) -> None:
    for col in range(2, last_col + 1):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    ws.cell(row=row, column=2, value=text).font = BOLD


def header_row(ws, row: int, first_col: int, values) -> None:
    for offset, value in enumerate(values):
        cell = ws.cell(row=row, column=first_col + offset, value=value)
        cell.font = HEADER
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def put(ws, row: int, col: int, value, font=FORMULA, fmt: str | None = None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    return cell


def col(index: int) -> str:
    return get_column_letter(index)


def absref(sheet: str, column: int, row: int) -> str:
    return f"'{sheet}'!${col(column)}${row}"


def define(wb, name: str, sheet: str, column: int, row: int) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=absref(sheet, column, row))
