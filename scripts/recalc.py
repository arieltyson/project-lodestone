"""Recalculate every formula in a workbook with headless LibreOffice and report errors.

openpyxl writes formulas without cached values, so anything that reads values (pandas,
the reconciliation test, GitHub's previewer) sees blanks until the file is recalculated.

Usage:
    python scripts/recalc.py model/lodestone_model.xlsx
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub RecalculateAndSave()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>"""

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")


def soffice() -> str:
    for candidate in ("soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if shutil.which(candidate) or Path(candidate).exists():
            return candidate
    raise SystemExit("LibreOffice (soffice) is required: brew install --cask libreoffice")


def recalc(path: Path, timeout: int = 180) -> dict:
    path = path.resolve()
    with tempfile.TemporaryDirectory(prefix="lodestone-lo-") as profile:
        profile_url = Path(profile).as_uri()
        subprocess.run([soffice(), "--headless", "--terminate_after_init",
                        f"-env:UserInstallation={profile_url}"], capture_output=True, timeout=timeout)
        macro_dir = Path(profile) / "user" / "basic" / "Standard"
        (macro_dir / "Module1.xba").write_text(MACRO)
        subprocess.run([soffice(), "--headless", "--norestore", f"-env:UserInstallation={profile_url}",
                        "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
                        str(path)], capture_output=True, timeout=timeout, check=True)
    return scan(path)


def scan(path: Path) -> dict:
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path)
    errors: dict[str, list[str]] = {}
    total = empty = 0
    for ws in formulas.worksheets:
        cached = values[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    total += 1
                    value = cached[cell.coordinate].value
                    if value is None:
                        empty += 1
                    elif isinstance(value, str) and value in ERRORS:
                        errors.setdefault(value, []).append(f"{ws.title}!{cell.coordinate}")
    return {
        "status": "errors_found" if errors or empty == total else "success",
        "total_formulas": total,
        "uncached_formulas": empty,
        "total_errors": sum(len(v) for v in errors.values()),
        "errors": {k: v[:25] for k, v in errors.items()},
    }


if __name__ == "__main__":
    result = recalc(Path(sys.argv[1]))
    print(json.dumps(result, indent=2))
    sys.exit(0 if result["status"] == "success" else 1)
