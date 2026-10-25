"""Build the peer trading comps table and write it into the model's Comps tab.

Writes data/comps/comps_table.csv (values, for diffs and the deck) and a Comps tab in
model/lodestone_model.xlsx whose multiples are live formulas on the Market tab, so a
refreshed snapshot flows through without rerunning this script.

Usage:
    python scripts/build_comps.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from lodestone.assumptions import Assumptions  # noqa: E402
from lodestone.comps import GRAMS_PER_OZ, LB_PER_TONNE, comps_table, peer_stats  # noqa: E402
from lodestone.workbook import styles as s  # noqa: E402
from lodestone.workbook.nav import INPUT_ROW0, MARKET_FIELDS, ROW_AU_DECK, ROW_CU_DECK  # noqa: E402

MODEL = ROOT / "model" / "lodestone_model.xlsx"
CSV = ROOT / "data" / "comps" / "comps_table.csv"
SHEET = "Comps"

COLUMNS = (
    ("Ticker", 10, None),
    ("Company", 20, None),
    ("Role", 9, None),
    ("Price (C$)", 10, s.USD_2),
    ("Mkt cap (US$M)", 12, s.USD_0),
    ("EV (US$M)", 12, s.USD_0),
    ("EV / EBITDA (TTM)", 11, s.MULT),
    ("P / CF (TTM)", 11, s.MULT),
    ("P / NAV", 9, s.MULT),
    ("Price / target", 10, s.MULT),
    ("P&P reserves (Mt CuEq)", 13, s.USD_2),
    ("EV / lb CuEq reserve (US$)", 13, s.USD_2),
)
STAT_METRICS = {"EV_EBITDA": 8, "PCF": 9, "PNAV": 10, "PTGT": 11, "EVLB": 13}


def market_col(field: str) -> str:
    keys = [f for f, _, _ in MARKET_FIELDS]
    return s.col(3 + keys.index(field))


def market_rows(wb) -> dict[str, int]:
    ws = wb["Market"]
    return {ws.cell(r, 2).value: r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value}


def reserves_formula(a: Assumptions, company: str) -> str:
    last = s.col(6 + a.last_year - a.first_year)
    ratio = f"(Inputs!${last}${ROW_AU_DECK}/(Inputs!${last}${ROW_CU_DECK}*{LB_PER_TONNE}))"
    terms = []
    for asset in a.assets_of(company):
        ore = f"('{asset}'!$C${INPUT_ROW0 + 1}-'{asset}'!$C${INPUT_ROW0 + 2})"
        terms.append(f"{ore}*'{asset}'!$C${INPUT_ROW0 + 3}+{ore}*'{asset}'!$C${INPUT_ROW0 + 4}/{GRAMS_PER_OZ}*{ratio}")
    return "=" + "+".join(terms)


def write_sheet(wb, a: Assumptions, table) -> None:
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET, index=wb.sheetnames.index("Corporate_NAV") + 1)
    s.title(ws, "Trading comparables",
            "Multiples are live formulas on the Market tab. Peer statistics exclude Centerra and Ero. "
            "Peer reserves are estimates (data/peer_inputs.csv); P/NAV shown only where a NAV is available.")
    s.header_row(ws, 4, 2, [c for c, _, _ in COLUMNS])
    for j, (_, width, _) in enumerate(COLUMNS):
        ws.column_dimensions[s.col(2 + j)].width = width

    rows = market_rows(wb)
    M = lambda field, r: f"Market!${market_col(field)}${r}"
    order = list(table.index)
    for i, ticker in enumerate(order):
        r, mr = 5 + i, rows[ticker]
        role = table.loc[ticker, "role"]
        company = a.corporate.index[a.corporate["ticker"] == ticker]
        values = [
            (ticker, s.BOLD),
            (f"={M('name', mr)}", s.LINK),
            (role, s.FORMULA),
            (f"={M('price_local', mr)}", s.LINK),
            (f"={M('market_cap_usd_m', mr)}", s.LINK),
            (f"={M('ev_usd_m', mr)}", s.LINK),
            (f"=IFERROR({M('ev_usd_m', mr)}/{M('ebitda_ttm_usd_m', mr)},\"n/a\")", s.FORMULA),
            (f"=IFERROR({M('market_cap_usd_m', mr)}/{M('cfo_ttm_usd_m', mr)},\"n/a\")", s.FORMULA),
        ]
        if len(company):
            values.append((f"={company[0]}_PNAV", s.LINK))
        else:
            navps = a._csv("peer_inputs.csv").set_index("ticker")["consensus_navps_cad"].get(ticker)
            values.append((f"={M('price_local', mr)}/{navps}" if navps == navps and navps else "n/a", s.FORMULA))
        values.append((f"=IFERROR({M('price_local', mr)}/{M('target_mean', mr)},\"n/a\")", s.FORMULA))
        if len(company):
            values.append((reserves_formula(a, company[0]), s.LINK))
        else:
            reserve = table.loc[ticker, "reserves_cueq_mt"]
            values.append((float(reserve) if reserve == reserve else "n/a", s.INPUT))
        values.append((f"=IFERROR(G{r}/(L{r}*{LB_PER_TONNE}),\"n/a\")", s.FORMULA))
        for j, (value, font) in enumerate(values):
            cell = s.put(ws, r, 2 + j, value, font, COLUMNS[j][2])
            if role != "peer":
                cell.fill = s.SECTION_FILL

    peer_rows = [5 + i for i, t in enumerate(order) if table.loc[t, "role"] == "peer"]
    first, last = min(peer_rows), max(peer_rows)
    r = last + 2
    for label, fn in (("Peer mean", "AVERAGE({})"), ("Peer median", "MEDIAN({})"),
                      ("Peer 25th percentile", "QUARTILE({},1)"), ("Peer 75th percentile", "QUARTILE({},3)")):
        s.put(ws, r, 2, label, s.BOLD)
        for name, c in STAT_METRICS.items():
            L = s.col(c)
            cell = s.put(ws, r, c, f"=IFERROR({fn.format(f'{L}{first}:{L}{last}')},\"n/a\")", s.BOLD, COLUMNS[c - 2][2])
            cell.border = s.TOTAL_BORDER
            suffix = {"Peer mean": "MEAN", "Peer median": "MED", "Peer 25th percentile": "Q1",
                      "Peer 75th percentile": "Q3"}[label]
            s.define(wb, f"COMPS_{name}_{suffix}", SHEET, c, r)
        r += 1
    r += 1
    s.put(ws, r, 2, "Notes", s.BOLD)
    for note in (
        "TTM EBITDA and operating cash flow from yfinance (latest reported); EV = market cap + debt - cash.",
        "P/NAV for Centerra and Ero links to the model's reserve-only NAV at the analyst deck (Corporate_NAV).",
        "Peer P/NAV needs a consensus NAV per share in data/peer_inputs.csv; none is published for free, so it shows n/a.",
        "Price / target uses the yfinance mean analyst target as a market-based cross-check.",
        "CuEq converts gold at the long-term deck ratio (gold price / copper price per tonne).",
    ):
        r += 1
        s.put(ws, r, 2, note, s.NOTE)


def main() -> int:
    a = Assumptions()
    table = comps_table(a)
    CSV.parent.mkdir(parents=True, exist_ok=True)
    table.round(4).to_csv(CSV)
    peer_stats(table).round(4).to_csv(CSV.with_name("peer_stats.csv"))
    wb = load_workbook(MODEL)
    write_sheet(wb, a, table)
    wb.save(MODEL)
    print(f"wrote {CSV.relative_to(ROOT)} and the {SHEET} tab of {MODEL.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
