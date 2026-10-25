"""Re-run the workbook under each price scenario and record the outputs.

Excel data tables do not survive openpyxl or LibreOffice, so the copper-price sensitivity
is produced the way an analyst would by hand: switch the price case on the Inputs tab,
recalculate, copy the outputs. Each run uses a scratch copy; the results are written into
the Scenarios tab of the real model as values.

Usage:
    python scripts/run_scenarios.py
"""

from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from lodestone.workbook.merger import SCEN_FIRST_ROW, SCENARIO_CASES, SCENARIO_OUTPUTS, SCENARIOS  # noqa: E402
from lodestone.workbook.nav import INPUTS, ROW_AU_CASE, ROW_CU_CASE, ROW_CU_CUSTOM  # noqa: E402
from scripts.recalc import recalc  # noqa: E402

MODEL = ROOT / "model" / "lodestone_model.xlsx"


def named_values(path: Path, names) -> dict[str, float]:
    wb = load_workbook(path, data_only=True)
    out = {}
    for name in names:
        sheet, coord = next(iter(wb.defined_names[name].destinations))
        out[name] = wb[sheet][coord.replace("$", "")].value
    return out


def run_case(cu_case: int, au_case: int, cu_custom: float | None, workdir: Path) -> dict[str, float]:
    scratch = workdir / "scenario.xlsx"
    shutil.copyfile(MODEL, scratch)
    wb = load_workbook(scratch)
    ws = wb[INPUTS]
    ws.cell(ROW_CU_CASE, 3).value = cu_case
    ws.cell(ROW_AU_CASE, 3).value = au_case
    if cu_custom is not None:
        ws.cell(ROW_CU_CUSTOM, 3).value = cu_custom
    wb.save(scratch)
    result = recalc(scratch)
    if result["status"] != "success":
        raise RuntimeError(f"recalc failed for case {cu_case}/{au_case}/{cu_custom}: {result}")
    return named_values(scratch, SCENARIO_OUTPUTS)


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="lodestone-scen-") as tmp:
        results = []
        for label, cu_case, au_case, cu in SCENARIO_CASES:
            results.append(run_case(cu_case, au_case, cu, Path(tmp)))
            print(f"{label:>14}: " + ", ".join(f"{k}={v:,.2f}" for k, v in results[-1].items()))
    wb = load_workbook(MODEL)
    ws = wb[SCENARIOS]
    for i, values in enumerate(results):
        for j, name in enumerate(SCENARIO_OUTPUTS):
            ws.cell(SCEN_FIRST_ROW + i, 6 + j).value = values[name]
    wb.save(MODEL)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
