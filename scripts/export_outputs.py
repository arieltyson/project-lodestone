"""Export the recalculated model's outputs to deck/model_outputs.json.

The pitch book, README and portfolio site read this file, so every figure they show comes
from the workbook rather than being typed in by hand.

Usage:
    python scripts/export_outputs.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from lodestone.assumptions import Assumptions  # noqa: E402
from lodestone.workbook.context import FIRST_YEAR_COL  # noqa: E402
from lodestone.workbook.merger import (  # noqa: E402
    CASH_GRID, COPPER_GRID, PREMIUM_GRID, SCEN_FIRST_ROW, SCENARIO_CASES, SCENARIO_OUTPUTS, SYNERGY_GRID,
)
from lodestone.workbook.nav import INPUT_ROW0, TS  # noqa: E402

MODEL = ROOT / "model" / "lodestone_model.xlsx"
OUTPUT = ROOT / "deck" / "model_outputs.json"


def table(ws, first_row: int, last_row: int, first_col: int, last_col: int) -> list[list]:
    return [[ws.cell(r, c).value for c in range(first_col, last_col + 1)] for r in range(first_row, last_row + 1)]


def find_row(ws, text: str, col: int = 2) -> int:
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, col).value == text:
            return r
    raise KeyError(f"{text!r} not found on {ws.title}")


def main() -> int:
    a = Assumptions()
    wb = load_workbook(MODEL, data_only=True)
    names = {}
    for name, dn in wb.defined_names.items():
        sheet, coord = next(iter(dn.destinations))
        names[name] = wb[sheet][coord.replace("$", "")].value

    years = a.years
    assets = {}
    for asset, row in a.assets.iterrows():
        ws = wb[asset]
        own = ws.cell(INPUT_ROW0, 3).value
        series = lambda key: [ws.cell(TS[key], FIRST_YEAR_COL + i).value * own for i in range(len(years))]
        assets[asset] = {
            "display_name": row["display_name"],
            "company": row["company"],
            "country": row["country"],
            "primary_metal": row["primary_metal"],
            "discount_rate": float(a.discount_rates[asset]),
            "nav": names[f"{asset}_NAV"],
            "life": names[f"{asset}_LIFE"],
            "aisc_y1": names[f"{asset}_AISC_Y1"],
            "aisc_lom": names[f"{asset}_AISC_LOM"],
            "reserve_mt": float(row["reserve_mt"]),
            "cu_grade_pct": float(row["cu_grade_pct"]),
            "au_gpt": float(row["au_gpt"]),
            "throughput_mtpa": float(row["throughput_mtpa"]),
            "expansion_mtpa": float(row["expansion_mtpa"]),
            "cu_mlb": series("cu_lb"),
            "au_koz": series("au_oz"),
            "ebitda": series("ebitda"),
            "fcf": series("fcf"),
        }

    comps_ws = wb["Comps"]
    comps_last = find_row(comps_ws, "Peer mean") - 2
    comps = table(comps_ws, 5, comps_last, 2, 13)
    comps_stats = {label: table(comps_ws, r, r, 2, 13)[0]
                   for label in ("Peer mean", "Peer median", "Peer 25th percentile", "Peer 75th percentile")
                   for r in [find_row(comps_ws, label)]}

    prec_ws = wb["Precedents"]
    prec_last = find_row(prec_ws, "Mean") - 2
    precedents = table(prec_ws, 5, prec_last, 2, 12)

    sens_ws = wb["Sensitivity"]
    sens_navps = table(sens_ws, 6, 5 + len(PREMIUM_GRID), 3, 2 + len(COPPER_GRID))
    eps_header = find_row(sens_ws, "Synergies \\ Cash %")
    sens_eps = table(sens_ws, eps_header + 1, eps_header + len(SYNERGY_GRID), 3, 2 + len(CASH_GRID))

    contrib_ws = wb["Contribution"]
    contribution = table(contrib_ws, 5, 12, 2, 6)

    scen_ws = wb["Scenarios"]
    scenarios = {
        label.split()[0] if cu is None else label: {
            output: scen_ws.cell(SCEN_FIRST_ROW + i, 6 + j).value for j, output in enumerate(SCENARIO_OUTPUTS)
        }
        for i, (label, _, _, cu) in enumerate(SCENARIO_CASES)
    }

    syn = a._csv("synergies.csv")
    out = {
        "timestamp_utc": a.macro["timestamp_utc"],
        "spot": {"copper": a.macro["copper_usd_lb"], "gold": a.macro["gold_usd_oz"], "cadusd": a.macro["cadusd"]},
        "deck": {"years": years, "copper": a.price_deck["copper_usd_lb"].tolist(),
                 "gold": a.price_deck["gold_usd_oz"].tolist()},
        "names": names,
        "assets": assets,
        "comps": comps,
        "comps_stats": comps_stats,
        "precedents": precedents,
        "sensitivity": {"premiums": PREMIUM_GRID, "copper": COPPER_GRID, "navps": sens_navps,
                        "synergies": SYNERGY_GRID, "cash": CASH_GRID, "eps_2028": sens_eps},
        "contribution": contribution,
        "scenarios": scenarios,
        "synergies": syn.to_dict(orient="records"),
        "market": {t: {k: (v.item() if hasattr(v, "item") else v) for k, v in row.items()}
                   for t, row in a.market.iterrows()},
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(out, indent=1, default=float, allow_nan=False))
    print(f"wrote {OUTPUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
