"""The Python engine must agree with the recalculated Excel model within 0.5%.

Reads cached values from model/lodestone_model.xlsx (written by `make model`), so this
runs anywhere without LibreOffice.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from lodestone.assumptions import COMPANIES, Assumptions
from lodestone.valuation import Prices, eps_sensitivity, football_field, navps_sensitivity, run
from lodestone.workbook.merger import (
    CASH_GRID, COPPER_GRID, PREMIUM_GRID, SCEN_FIRST_ROW, SCENARIO_OUTPUTS, SYNERGY_GRID,
)

MODEL = Path(__file__).resolve().parents[1] / "model" / "lodestone_model.xlsx"
TOLERANCE = 0.005


def close(python: float, excel: float) -> bool:
    return python == pytest.approx(excel, rel=TOLERANCE, abs=1e-6)


@pytest.fixture(scope="module")
def wb():
    return load_workbook(MODEL, data_only=True)


@pytest.fixture(scope="module")
def excel(wb):
    values = {}
    for name, dn in wb.defined_names.items():
        sheet, coord = next(iter(dn.destinations))
        values[name] = wb[sheet][coord.replace("$", "")].value
    return values


@pytest.fixture(scope="module")
def a():
    return Assumptions()


@pytest.fixture(scope="module")
def model(a):
    return run(a)


def test_workbook_is_recalculated(excel):
    assert excel["CG_NAV"] is not None, "run `make model` so the workbook has cached values"


@pytest.mark.parametrize("asset", ["MtMilligan", "Oksut", "Goldfield", "Caraiba", "Tucuma", "Xavantina"])
@pytest.mark.parametrize("field", ["npv", "nav", "life", "aisc_y1", "aisc_lom"])
def test_asset_outputs(model, excel, asset, field):
    company = next(c for c in model.companies.values() if asset in c.assets)
    assert close(getattr(company.assets[asset], field), excel[f"{asset}_{field.upper()}"])


@pytest.mark.parametrize("company", COMPANIES)
@pytest.mark.parametrize("field", ["asset_nav", "ga", "cash", "debt", "nav", "shares", "navps_usd", "navps_cad",
                                   "pnav"])
def test_corporate_nav(model, excel, company, field):
    assert close(getattr(model.companies[company], field), excel[f"{company}_{field.upper()}"])


@pytest.mark.parametrize("company", COMPANIES)
@pytest.mark.parametrize("year", [2027, 2028])
@pytest.mark.parametrize("field", ["ebitda", "ni", "eps", "cfo", "cfps"])
def test_financials(model, excel, company, year, field):
    assert close(model.companies[company].fin[year][field], excel[f"{company}_{field.upper()}_{year}"])


DEAL_FIELDS = {
    "offer_cad": "DEAL_OFFER_CAD", "exchange_ratio": "DEAL_ER", "new_shares": "DEAL_NEW_SHARES",
    "equity_usd": "DEAL_EQUITY_USD", "cash_usd": "DEAL_CASH_USD", "stock_usd": "DEAL_STOCK_USD",
    "net_debt_assumed": "DEAL_NET_DEBT", "ev_usd": "DEAL_EV_USD", "pnav": "DEAL_PNAV", "ev_ebitda": "DEAL_EV_EBITDA",
    "cash_need": "DEAL_CASH_NEED", "cash_used": "DEAL_CASH_USED", "new_debt": "DEAL_NEW_DEBT",
    "pf_shares": "PF_SHARES", "own_cg": "PF_OWN_CG", "syn_npv": "SYN_NPV", "pf_nav": "PF_NAV",
    "pf_navps": "PF_NAVPS", "acc_navps": "ACC_NAVPS", "acc_navps_ex_syn": "ACC_NAVPS_EX_SYN",
    "pf_net_debt": "PF_NET_DEBT", "leverage": "PF_LEVERAGE",
}


@pytest.mark.parametrize("field,name", DEAL_FIELDS.items())
def test_deal(model, excel, field, name):
    assert close(getattr(model.deal, field), excel[name])


@pytest.mark.parametrize("year", [2027, 2028])
@pytest.mark.parametrize("field,name", [("pf_ni", "PF_NI"), ("pf_eps", "PF_EPS"), ("acc_eps", "ACC_EPS"),
                                        ("pf_cfps", "PF_CFPS"), ("acc_cfps", "ACC_CFPS"), ("pf_ebitda", "PF_EBITDA")])
def test_accretion(model, excel, year, field, name):
    assert close(model.deal.years[year][field], excel[f"{name}_{year}"])


def test_spot_scenario(a, wb):
    spot = run(a, Prices.spot(a))
    ws = wb["Scenarios"]
    row = SCEN_FIRST_ROW + 1
    for j, name in enumerate(SCENARIO_OUTPUTS):
        company, field = name.split("_", 1)
        assert close(getattr(spot.companies[company], field.lower()), ws.cell(row, 6 + j).value), name


def test_navps_sensitivity(a, wb):
    grid = navps_sensitivity(a, PREMIUM_GRID, COPPER_GRID)
    ws = wb["Sensitivity"]
    for i, prem in enumerate(PREMIUM_GRID):
        for j, cu in enumerate(COPPER_GRID):
            assert close(grid.loc[prem, cu], ws.cell(6 + i, 3 + j).value), (prem, cu)


def test_eps_sensitivity(model, wb):
    grid = eps_sensitivity(model, SYNERGY_GRID, CASH_GRID)
    ws = wb["Sensitivity"]
    first = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "Synergies \\ Cash %") + 1
    for i, syn in enumerate(SYNERGY_GRID):
        for j, c in enumerate(CASH_GRID):
            assert close(grid.loc[syn, c], ws.cell(first + i, 3 + j).value), (syn, c)


def test_football_field(a, model, excel):
    ff = football_field(a, model)
    for method, row in ff.iterrows():
        assert close(row["low"], excel[f"FF_{method}_LOW"]), method
        assert close(row["high"], excel[f"FF_{method}_HIGH"]), method
